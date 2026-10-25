from __future__ import annotations

import csv
import hashlib
import json
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .config import DATE_FIELDS, FIELD_SPECS, MAX_DATA_ROWS, MAX_UPLOAD_BYTES, MONEY_FIELDS

CENT = Decimal("0.01")


class ImportProblem(ValueError):
    pass


@dataclass
class ParsedFile:
    headers: list[str]
    rows: list[list[Any]]
    sheet_names: list[str]
    selected_sheet: str


@dataclass
class CleanRow:
    row_number: int
    payload: dict[str, Any]
    data: dict[str, Any]
    row_hash: str
    business_key: str
    errors: list[str]
    status: str = "accepted"


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def _unique_headers(values: list[Any]) -> list[str]:
    values = list(values)
    while values and (values[-1] is None or not str(values[-1]).strip()):
        values.pop()
    result: list[str] = []
    counts: dict[str, int] = {}
    for i, value in enumerate(values, start=1):
        base = _cell_text(value) or f"未命名列{i}"
        counts[base] = counts.get(base, 0) + 1
        result.append(base if counts[base] == 1 else f"{base}（{counts[base]}）")
    return result


def parse_upload(path: str | Path, filename: str, sheet_name: str | None = None) -> ParsedFile:
    p = Path(path)
    ext = Path(filename).suffix.lower()
    size = p.stat().st_size
    if size == 0:
        raise ImportProblem("文件为空。")
    if size > MAX_UPLOAD_BYTES:
        raise ImportProblem("文件超过 20 MB 限制。")
    if ext not in {".xlsx", ".csv", ".tsv"}:
        raise ImportProblem("仅支持 .xlsx、.csv 和 .tsv；旧版 .xls 请先另存为 .xlsx。")

    if ext == ".xlsx":
        try:
            wb = load_workbook(p, read_only=True, data_only=True)
            names = list(wb.sheetnames)
            chosen = sheet_name if sheet_name in names else names[0]
            ws = wb[chosen]
            iterator = ws.iter_rows(values_only=True)
            first = next(iterator, None)
            if first is None:
                raise ImportProblem("所选工作表没有内容。")
            headers = _unique_headers(list(first))
            rows: list[list[Any]] = []
            for row in iterator:
                if all(v is None or str(v).strip() == "" for v in row):
                    continue
                rows.append(list(row[: len(headers)]))
                if len(rows) > MAX_DATA_ROWS:
                    raise ImportProblem("单个工作表超过 50,000 行限制。")
            wb.close()
            return ParsedFile(headers, rows, names, chosen)
        except ImportProblem:
            raise
        except Exception as exc:
            raise ImportProblem(f"无法读取 xlsx：{exc}") from exc

    data: bytes = p.read_bytes()
    decoded = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            decoded = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise ImportProblem("CSV 编码无法识别，请另存为 UTF-8 或 GB18030。")
    try:
        sample = decoded[:8192]
        delimiter = "\t" if ext == ".tsv" else csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = "\t" if ext == ".tsv" else ","
    reader = csv.reader(decoded.splitlines(), delimiter=delimiter)
    first = next(reader, None)
    if not first:
        raise ImportProblem("CSV 没有表头。")
    headers = _unique_headers(first)
    rows = []
    for row in reader:
        if not row or all(not str(v).strip() for v in row):
            continue
        rows.append(list(row[: len(headers)]) + [None] * max(0, len(headers) - len(row)))
        if len(rows) > MAX_DATA_ROWS:
            raise ImportProblem("文件超过 50,000 行限制。")
    return ParsedFile(headers, rows, ["CSV"], "CSV")


def _fold(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).strip().lower()
    return re.sub(r"[\s_\-（）()【】\[\]：:,.，。/\\]+", "", text)


def suggest_mapping(report_type: str, headers: list[str]) -> dict[str, str]:
    folded = {_fold(h): h for h in headers}
    result: dict[str, str] = {}
    for key, _label, _required, aliases in FIELD_SPECS[report_type]:
        candidates = [key, *aliases]
        found = next((folded[_fold(alias)] for alias in candidates if _fold(alias) in folded), "")
        result[key] = found
    return result


def _decimal(value: Any, field: str, allow_negative: bool = False) -> Decimal:
    if value is None or str(value).strip() == "":
        raise ImportProblem(f"{field}为空")
    if isinstance(value, bool):
        raise ImportProblem(f"{field}不是有效数字")
    text = str(value).strip().replace(",", "").replace("，", "").replace("￥", "").replace("¥", "").replace("元", "")
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError) as exc:
        raise ImportProblem(f"{field}不是有效数字：{value}") from exc
    if not number.is_finite():
        raise ImportProblem(f"{field}不是有限数值")
    if negative:
        number = -abs(number)
    if not allow_negative and number < 0:
        raise ImportProblem(f"{field}不能为负数")
    return number


def money_fen(value: Any, field: str, allow_negative: bool = False) -> int:
    amount = _decimal(value, field, allow_negative).quantize(CENT, rounding=ROUND_HALF_UP)
    return int(amount * 100)


def _date(value: Any, field: str) -> str:
    if value is None or str(value).strip() == "":
        raise ImportProblem(f"{field}为空")
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return from_excel(value).date().isoformat()
        except Exception as exc:
            raise ImportProblem(f"{field}不是有效日期") from exc
    text = str(value).strip()
    text = re.sub(r"[年/.]", "-", text).replace("月", "-").replace("日", "")
    text = re.sub(r"\s*-\s*", "-", text)
    text = text.replace("T", " ")
    text = text.strip()
    try:
        return datetime.fromisoformat(text).date().isoformat()
    except ValueError:
        for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y%m%d", "%Y-%m-%d %H:%M"):
            try:
                return datetime.strptime(text, fmt).date().isoformat()
            except ValueError:
                pass
    raise ImportProblem(f"{field}不是有效日期：{value}")


def _status_class(value: str) -> str:
    v = _fold(value)
    if any(word in v for word in ("取消", "关闭", "closed", "cancel", "失败", "作废")):
        return "cancelled"
    if any(word in v for word in ("完成", "成功", "已支付", "已付款", "已发货", "待收货", "待评价", "交易成功", "finished", "paid", "completed", "delivered")):
        return "valid"
    return "unknown"


def _refund_confirmed(value: str) -> bool:
    v = _fold(value)
    if any(word in v for word in ("撤销", "拒绝", "关闭", "失败", "cancel", "refus", "fail")):
        return False
    return any(word in v for word in ("成功", "完成", "已退款", "confirmed", "success", "completed", "refunded"))


def _settlement_confirmed(value: str) -> bool | None:
    v = _fold(value)
    if any(word in v for word in ("待", "处理中", "未结算", "冻结", "失败", "撤销", "pending", "processing", "fail", "cancel")):
        return False
    if any(word in v for word in ("已结算", "已到账", "成功", "完成", "settled", "completed", "success")):
        return True
    return None


def _json_value(v: Any) -> Any:
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    return str(v)


def _hash(value: Any) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")).hexdigest()


def _business_key(report_type: str, data: dict[str, Any], row_hash: str) -> str:
    if report_type == "orders":
        line = data.get("platform_item_id") or data["sku_id"]
        return f"order:{data['platform_order_id']}:{line}"
    if report_type == "refunds":
        return f"refund:{data['platform_refund_id']}" if data.get("platform_refund_id") else f"refund:{data['platform_order_id']}:{data.get('platform_item_id') or ''}:{data['completed_at']}:{data['refund_amount_fen']}"
    if report_type == "fees":
        return f"fee:{data['fee_id']}" if data.get("fee_id") else f"fee:{data.get('platform_order_id') or ''}:{data['fee_date']}:{data['fee_type']}:{data['fee_amount_fen']}"
    if report_type == "settlements":
        return f"settlement:{data['transaction_id']}" if data.get("transaction_id") else f"settlement:{data.get('platform_order_id') or ''}:{data['settlement_date']}:{data['settlement_amount_fen']}"
    if report_type == "costs":
        return f"cost:{data['sku_id']}:{data['effective_from']}"
    if report_type == "expenses":
        return f"expense:{data['expense_id']}" if data.get("expense_id") else f"expense:{data.get('platform_order_id') or ''}:{data.get('sku_id') or ''}:{data['expense_date']}:{data['expense_type']}:{data['expense_amount_fen']}"
    return row_hash


def clean_rows(report_type: str, headers: list[str], rows: list[list[Any]], mapping: dict[str, str]) -> list[CleanRow]:
    if report_type not in FIELD_SPECS:
        raise ImportProblem("不支持的报表类型")
    mapped_required = []
    for key, label, required, _aliases in FIELD_SPECS[report_type]:
        if required and not mapping.get(key):
            mapped_required.append(label)
    if report_type == "orders" and not mapping.get("item_amount") and not mapping.get("unit_price"):
        mapped_required.append("商品销售金额或商品单价")
    if mapped_required:
        raise ImportProblem("请先映射必填字段：" + "、".join(mapped_required))
    indexes = {key: headers.index(h) for key, h in mapping.items() if h in headers}
    result: list[CleanRow] = []
    for row_number, row in enumerate(rows, start=2):
        payload = {key: _json_value(row[idx] if idx < len(row) else None) for key, idx in indexes.items()}
        if not any(v not in (None, "") for v in payload.values()):
            continue
        errors: list[str] = []
        data: dict[str, Any] = {}
        try:
            def raw(key: str) -> Any:
                return payload.get(key)

            if report_type == "orders":
                data = {
                    "platform_order_id": _cell_text(raw("platform_order_id")),
                    "ordered_at": _date(raw("ordered_at"), "下单日期"),
                    "order_status": _cell_text(raw("order_status")),
                    "status_class": _status_class(_cell_text(raw("order_status"))),
                    "sku_id": _cell_text(raw("sku_id")),
                    "product_name": _cell_text(raw("product_name")) or None,
                    "quantity": str(_decimal(raw("quantity"), "销售数量")),
                    "platform_item_id": _cell_text(raw("platform_item_id")) or None,
                    "discount_fen": money_fen(raw("discount_amount") or 0, "商品优惠金额"),
                    "unit_price_fen": money_fen(raw("unit_price"), "商品单价") if raw("unit_price") not in (None, "") else None,
                }
                if raw("item_amount") not in (None, ""):
                    data["item_amount_fen"] = money_fen(raw("item_amount"), "商品销售金额")
                else:
                    qty = _decimal(raw("quantity"), "销售数量")
                    unit = _decimal(raw("unit_price"), "商品单价")
                    amount = (qty * unit - Decimal(data["discount_fen"]) / 100).quantize(CENT, rounding=ROUND_HALF_UP)
                    if amount < 0:
                        raise ImportProblem("优惠金额不能大于商品金额")
                    data["item_amount_fen"] = int(amount * 100)
                if not data["platform_order_id"] or not data["sku_id"]:
                    raise ImportProblem("订单号和 SKU 编号不能为空")
                if _decimal(raw("quantity"), "销售数量") <= 0:
                    raise ImportProblem("销售数量必须大于 0")
                if not data["order_status"]:
                    raise ImportProblem("订单状态为空，不能判断是否计入经营指标")
                if data["status_class"] == "unknown":
                    errors.append("订单状态无法归类：该行不计入销售统计")
            elif report_type == "refunds":
                data = {
                    "platform_refund_id": _cell_text(raw("platform_refund_id")) or None,
                    "platform_order_id": _cell_text(raw("platform_order_id")),
                    "platform_item_id": _cell_text(raw("platform_item_id")) or None,
                    "refund_status": _cell_text(raw("refund_status")),
                    "confirmed": _refund_confirmed(_cell_text(raw("refund_status"))),
                    "refund_amount_fen": money_fen(raw("refund_amount"), "退款金额"),
                    "completed_at": _date(raw("completed_at"), "退款完成日期"),
                }
                if not data["platform_order_id"]:
                    raise ImportProblem("平台订单号不能为空")
                if not data["refund_status"]:
                    raise ImportProblem("退款状态为空")
                if not data["confirmed"]:
                    errors.append("退款尚未确认完成：该行不冲减已确认退款金额")
            elif report_type == "fees":
                data = {
                    "fee_id": _cell_text(raw("fee_id")) or None,
                    "platform_order_id": _cell_text(raw("platform_order_id")) or None,
                    "fee_type": _cell_text(raw("fee_type")),
                    "fee_amount_fen": money_fen(raw("fee_amount"), "费用金额", allow_negative=True),
                    "fee_date": _date(raw("fee_date"), "费用发生日期"),
                }
                if not data["fee_type"]:
                    raise ImportProblem("费用类型不能为空")
            elif report_type == "settlements":
                settlement_status = _cell_text(raw("settlement_status"))
                is_confirmed = _settlement_confirmed(settlement_status)
                data = {
                    "transaction_id": _cell_text(raw("transaction_id")) or None,
                    "platform_order_id": _cell_text(raw("platform_order_id")) or None,
                    "settlement_amount_fen": money_fen(raw("settlement_amount"), "结算金额", allow_negative=True),
                    "settlement_status": settlement_status,
                    "is_confirmed": bool(is_confirmed),
                    "settlement_date": _date(raw("settlement_date"), "结算日期"),
                }
                if not settlement_status:
                    raise ImportProblem("结算状态为空")
                if is_confirmed is None:
                    errors.append("结算状态无法识别：该行不计入已确认结算金额")
                    data["is_confirmed"] = False
                elif is_confirmed is False:
                    errors.append("结算尚未确认：该行不计入已确认结算金额")
            elif report_type == "costs":
                effective_from = _date(raw("effective_from"), "成本生效日期") if raw("effective_from") not in (None, "") else "1900-01-01"
                effective_to = _date(raw("effective_to"), "成本失效日期") if raw("effective_to") not in (None, "") else None
                if effective_to and effective_to < effective_from:
                    raise ImportProblem("成本失效日期早于生效日期")
                data = {
                    "sku_id": _cell_text(raw("sku_id")),
                    "product_name": _cell_text(raw("product_name")) or None,
                    "unit_cost_fen": money_fen(raw("unit_cost"), "单位成本"),
                    "effective_from": effective_from,
                    "effective_to": effective_to,
                }
                if not data["sku_id"]:
                    raise ImportProblem("SKU 编号不能为空")
            elif report_type == "expenses":
                data = {
                    "expense_id": _cell_text(raw("expense_id")) or None,
                    "expense_type": _cell_text(raw("expense_type")),
                    "expense_amount_fen": money_fen(raw("expense_amount"), "费用金额"),
                    "expense_date": _date(raw("expense_date"), "费用发生日期"),
                    "allocation_method": _cell_text(raw("allocation_method")).lower() or "store",
                    "platform_order_id": _cell_text(raw("platform_order_id")) or None,
                    "sku_id": _cell_text(raw("sku_id")) or None,
                    "note": _cell_text(raw("note")) or None,
                }
                if not data["expense_type"]:
                    raise ImportProblem("费用类型不能为空")
                if data["allocation_method"] not in {"store", "店铺", "order", "订单", "sku", "商品"}:
                    raise ImportProblem("分摊方式只接受 store/店铺、order/订单、sku/商品")
                data["allocation_method"] = {"店铺": "store", "订单": "order", "商品": "sku"}.get(data["allocation_method"], data["allocation_method"])
                if data["allocation_method"] == "order" and not data["platform_order_id"]:
                    raise ImportProblem("订单分摊费用必须填写关联订单号")
                if data["allocation_method"] == "sku" and not data["sku_id"]:
                    raise ImportProblem("商品分摊费用必须填写关联 SKU")
        except ImportProblem as exc:
            data = {}
            errors.append(str(exc))
        row_hash = _hash(data if data else payload)
        business_key = _business_key(report_type, data, row_hash) if data else f"invalid:{row_hash}"
        result.append(CleanRow(row_number, payload, data, row_hash, business_key, errors,
                               status="accepted" if data else "rejected"))
    return result


def classify_duplicates(rows: list[CleanRow], existing: dict[str, str]) -> None:
    seen: dict[str, str] = {}
    for row in rows:
        if not row.data:
            row.status = "rejected"
            continue
        if row.business_key in seen:
            row.status = "duplicate" if seen[row.business_key] == row.row_hash else "conflict"
            row.errors.append("本文件内的完全重复行，已跳过" if row.status == "duplicate" else "本文件内业务键相同但内容不同，待人工复核")
            continue
        seen[row.business_key] = row.row_hash
        if row.business_key in existing:
            row.status = "duplicate" if existing[row.business_key] == row.row_hash else "conflict"
            row.errors.append("与已导入记录完全重复，已跳过" if row.status == "duplicate" else "与已导入记录业务键相同但内容不同，待人工复核")


def normalize_row(row: CleanRow) -> dict[str, Any]:
    return row.data
