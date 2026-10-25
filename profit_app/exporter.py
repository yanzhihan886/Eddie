from __future__ import annotations

from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .calculations import build_summary, fmt_money, product_analysis

NAVY = "17365D"
BLUE = "285E93"
PALE = "EAF1F8"
GRAY = "F4F7FA"
WHITE = "FFFFFF"
ORANGE = "FCE4D6"
THIN_BLUE = Side(style="thin", color="B9C9D9")


def _safe_text(value: Any) -> Any:
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _write_sheet(wb: Workbook, name: str, headers: list[str], rows: list[list[Any]], widths: list[int] | None = None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Microsoft YaHei", color=WHITE, bold=True, size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for r_idx, row in enumerate(rows, start=2):
        ws.append([_safe_text(v) for v in row])
        for cell in ws[r_idx]:
            cell.font = Font(name="Microsoft YaHei", size=10, color="243746")
            cell.fill = PatternFill("solid", fgColor=WHITE if r_idx % 2 == 0 else GRAY)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN_BLUE)
    if widths:
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
    else:
        for i, header in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(14, len(str(header)) + 3), 30)
    ws.auto_filter.ref = ws.dimensions if ws.max_row > 1 else f"A1:{get_column_letter(len(headers))}1"
    return ws


def create_report(db, store_id: str | None, platform: str | None, start: str, end: str) -> BytesIO:
    summary = build_summary(db, store_id, platform, start, end)
    wb = Workbook()
    overview = wb.active
    overview.title = "经营汇总"
    overview.sheet_view.showGridLines = False
    overview.merge_cells("A1:D1")
    overview["A1"] = "卖家利润助手｜本地经营分析报告"
    overview["A1"].fill = PatternFill("solid", fgColor=NAVY)
    overview["A1"].font = Font(name="Microsoft YaHei", color=WHITE, size=16, bold=True)
    overview["A1"].alignment = Alignment(vertical="center")
    overview.row_dimensions[1].height = 34
    generated = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    filters = f"店铺={store_id or '全部'}；平台={platform or '全部'}"
    metadata = [
        ["统计期间", f"{start} 至 {end}"],
        ["筛选条件", filters],
        ["生成时间", generated],
        ["利润规则版本", "event-period-fen-v1"],
        ["数据来源", "用户上传的本地文件；平台官方格式验证状态见下方说明"],
    ]
    for row in metadata:
        overview.append(row)
    overview.append(["指标", "金额/数值", "说明", "数据状态"])
    values = [
        ("商品销售额", summary["gross_sales_fen"], "已支付/已完成订单按下单日期", "已导入数据"),
        ("确认退款", summary["refund_fen"], "按退款完成日期，不要求订单日期同月", "仅已确认退款"),
        ("净销售额", summary["net_sales_fen"], "商品销售额 − 本期确认退款", "事件期间口径"),
        ("已匹配商品成本", summary["known_cost_fen"], "SKU 成本按订单日期有效版本匹配；不自动冲回退货成本", "缺成本 SKU 不计入此项"),
        ("已导入平台费用", summary["fee_fen"], "按费用发生日期；缺少账单不当作零", "只包含已导入记录"),
        ("已导入其他经营费用", summary["expense_fen"], "按费用发生日期；不含未提供的费用", "只包含已导入记录"),
        ("已知项目利润估算", summary["known_profit_fen"], "净销售额 − 已匹配成本 − 已导入平台费用 − 已导入其他费用", summary["completeness"]["label"]),
        ("利润率（已知项目口径）", f"{summary['profit_margin']:.2%}" if summary["profit_margin"] is not None else "不可计算", "已知项目利润估算 ÷ 净销售额", summary["completeness"]["label"]),
        ("订单数", summary["order_count"], "去重后的有效订单号数量", "已支付/已完成"),
        ("销售件数", str(summary["item_count"]), "有效订单的商品数量", "已导入数据"),
        ("取消/关闭订单行", summary["cancelled_count"], "不计入销售额", "已排除"),
        ("结算金额", summary["settlement_fen"], "按结算日期，仅为导入记录汇总", "不等同于销售额减退款"),
        ("未解决异常数", summary["open_issue_count"], "请在异常清单中复核", "包含待复核/打开项"),
    ]
    for metric, value, note, status in values:
        if isinstance(value, int) and metric not in ("订单数", "销售件数", "取消/关闭订单行", "未解决异常数"):
            value = value / 100
        overview.append([metric, value, note, status])
    for cell in overview[7]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Microsoft YaHei", color=WHITE, bold=True)
    money_metrics = {"商品销售额", "确认退款", "净销售额", "已匹配商品成本", "已导入平台费用", "已导入其他经营费用", "已知项目利润估算", "结算金额"}
    for row in overview.iter_rows(min_row=8):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color="243746")
            cell.fill = PatternFill("solid", fgColor=WHITE if cell.row % 2 == 0 else GRAY)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN_BLUE)
        if row[0].value in money_metrics:
            row[1].number_format = '¥#,##0.00;[Red]-¥#,##0.00'
    note_row = overview.max_row + 2
    overview.cell(note_row, 1, "重要口径与限制")
    overview.cell(note_row, 1).font = Font(name="Microsoft YaHei", bold=True, color=NAVY)
    limitations = [
        "本报告为经营分析估算，不是会计记账或纳税申报结果。缺失成本/费用不视为零；利润数字仅为已知项目口径。",
        "订单按下单日期、退款按退款完成日期、平台费用/其他费用按发生日期、结算按结算日期统计；跨月事件各按自身日期归属。",
        "退款默认不冲回商品成本。商品级利润中的订单级退款/费用按商品金额比例分摊；店铺级费用不强行分摊给商品。",
        "订单级结算差异只对齐已关联订单、费用和结算记录，以销售−退款−费用作简化复核线索，不代表平台官方账单结论。",
        "淘宝/天猫、拼多多、抖音电商当前均使用通用字段映射；没有真实脱敏平台报表样本验证，不能视为已适配官方导出格式。",
    ]
    for i, text in enumerate(limitations, start=note_row + 1):
        overview.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
        overview.cell(i, 1, _safe_text(text))
        overview.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
        overview.cell(i, 1).fill = PatternFill("solid", fgColor=ORANGE if i == note_row + 1 else PALE)
        overview.row_dimensions[i].height = 32
    for col, width in enumerate((30, 22, 66, 32), 1):
        overview.column_dimensions[get_column_letter(col)].width = width
    overview.freeze_panes = "A7"

    ids = summary["store_ids"]
    marks = ",".join("?" for _ in ids) or "NULL"
    p = [*ids, start, end]
    order_rows = db.execute(f"SELECT orders.*,stores.name AS store_name,stores.platform FROM orders JOIN stores ON stores.id=orders.store_id WHERE orders.store_id IN ({marks}) AND ordered_at BETWEEN ? AND ? ORDER BY ordered_at,platform_order_id", p).fetchall()
    order_data = []
    for r in order_rows:
        cost = db.execute("SELECT unit_cost_fen FROM product_costs WHERE store_id=? AND sku_id=? AND effective_from<=? AND (effective_to IS NULL OR effective_to>=?) ORDER BY effective_from DESC LIMIT 1", (r["store_id"], r["sku_id"], r["ordered_at"], r["ordered_at"])).fetchone()
        line_cost = int((Decimal(cost["unit_cost_fen"]) * Decimal(str(r["quantity"]))).quantize(Decimal("1"), rounding=ROUND_HALF_UP)) if cost else None
        order_data.append([r["store_name"], r["platform"], r["platform_order_id"], r["platform_item_id"], r["ordered_at"], r["order_status"], r["sku_id"], r["product_name"], r["quantity"], r["item_amount_fen"] / 100, line_cost / 100 if line_cost is not None else "缺失", r["source_record_id"]])
    ws = _write_sheet(wb, "订单明细", ["店铺", "平台", "订单号", "明细号", "下单日期", "状态", "SKU", "商品名称", "数量", "商品销售额", "匹配成本", "原始记录ID"], order_data, [22, 16, 24, 20, 14, 15, 18, 28, 12, 16, 16, 38])
    for row in ws.iter_rows(min_row=2, min_col=10, max_col=11):
        row[0].number_format = row[1].number_format = '¥#,##0.00;[Red]-¥#,##0.00'

    refund_rows = db.execute(f"SELECT refunds.*,stores.name AS store_name,stores.platform FROM refunds JOIN stores ON stores.id=refunds.store_id WHERE refunds.store_id IN ({marks}) AND completed_at BETWEEN ? AND ? ORDER BY completed_at", p).fetchall()
    _write_sheet(wb, "退款明细", ["店铺", "平台", "退款单号", "订单号", "明细号", "状态", "已确认", "退款完成日期", "退款金额", "来源记录ID"],
                 [[r["store_name"], r["platform"], r["platform_refund_id"], r["platform_order_id"], r["platform_item_id"], r["refund_status"], "是" if r["is_confirmed"] else "否", r["completed_at"], r["refund_amount_fen"]/100, r["source_record_id"]] for r in refund_rows],
                 [22, 16, 22, 22, 20, 16, 12, 16, 16, 38])
    fee_rows = db.execute(f"SELECT fees.*,stores.name AS store_name,stores.platform FROM fees JOIN stores ON stores.id=fees.store_id WHERE fees.store_id IN ({marks}) AND fee_date BETWEEN ? AND ? ORDER BY fee_date", p).fetchall()
    _write_sheet(wb, "平台费用", ["店铺", "平台", "费用流水", "订单号", "费用类型", "费用日期", "金额", "来源记录ID"],
                 [[r["store_name"], r["platform"], r["fee_id"], r["platform_order_id"], r["fee_type"], r["fee_date"], r["fee_amount_fen"]/100, r["source_record_id"]] for r in fee_rows],
                 [22, 16, 22, 22, 22, 16, 16, 38])
    products, refund_unallocated, fee_unallocated, expense_unallocated = product_analysis(db, summary, start, end)
    _write_sheet(wb, "商品成本利润", ["平台/店铺内SKU", "SKU", "商品名称", "销量", "商品销售额", "分摊退款", "已匹配成本", "订单级费用", "直接分摊其他费用", "已知项目利润", "成本缺失"],
                 [["跨店汇总", x["sku_id"], x["product_name"], str(x["quantity"]), x["sales_fen"]/100, x["refund_fen"]/100, x["cost_fen"]/100, x["fee_fen"]/100, x["expense_fen"]/100, x["known_profit_fen"]/100, "是" if x["missing_cost"] else "否"] for x in products],
                 [24, 20, 28, 14, 16, 16, 16, 16, 20, 18, 14])
    ws = wb["商品成本利润"]
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=10):
        for cell in row:
            cell.number_format = '¥#,##0.00;[Red]-¥#,##0.00'
    ws.append(["未分配退款", "", "无法匹配商品或跨期订单不在商品销售期间", "", "", refund_unallocated/100])
    ws.append(["未分配平台费用", "", "订单未在当前商品销售期间，保留在店铺汇总", "", "", "", "", fee_unallocated/100])
    ws.append(["未分配其他费用", "", "店铺级费用不会强行摊到 SKU", "", "", "", "", "", expense_unallocated/100])

    settlements = db.execute(f"SELECT settlements.*,stores.name AS store_name,stores.platform FROM settlements JOIN stores ON stores.id=settlements.store_id WHERE settlements.store_id IN ({marks}) AND settlement_date BETWEEN ? AND ?", p).fetchall()
    issues = db.execute(f"SELECT issues.*,stores.name AS store_name FROM issues JOIN stores ON stores.id=issues.store_id WHERE issues.store_id IN ({marks}) AND issue_type='settlement_mismatch' ORDER BY created_at DESC", ids).fetchall() if ids else []
    diff_rows = [[r["store_name"], r["issue_type"], r["related_record_ids"], r["amount_difference_fen"]/100 if r["amount_difference_fen"] is not None else None, r["status"], r["description"], r["resolution_note"]] for r in issues]
    if not diff_rows:
        diff_rows = [["—", "无订单级简化差异", "", "", "—", "注意：没有差异并不证明结算完全匹配；未关联订单或汇总账单不会自动推断。", ""]]
    _write_sheet(wb, "结算差异", ["店铺", "异常类型", "关联记录", "差异金额", "状态", "说明", "处理备注"], diff_rows, [22, 24, 26, 16, 14, 80, 38])
    _write_sheet(wb, "数据缺失与异常", ["店铺", "平台", "异常类型", "严重程度", "状态", "差异金额", "描述", "处理备注", "来源记录ID"],
                 [[r["store_name"], r["platform"], r["issue_type"], r["severity"], r["status"], r["amount_difference_fen"]/100 if r["amount_difference_fen"] is not None else "", r["description"], r["resolution_note"], r["source_record_id"]] for r in db.execute(f"SELECT issues.*,stores.name AS store_name,stores.platform FROM issues JOIN stores ON stores.id=issues.store_id WHERE issues.store_id IN ({marks}) ORDER BY severity,created_at DESC", ids)] if ids else [],
                 [22, 16, 26, 14, 14, 16, 80, 38, 38])
    _write_sheet(wb, "结算记录", ["店铺", "平台", "流水号", "订单号", "状态", "结算日期", "结算金额", "来源记录ID"],
                 [[r["store_name"], r["platform"], r["transaction_id"], r["platform_order_id"], r["settlement_status"], r["settlement_date"], r["settlement_amount_fen"]/100, r["source_record_id"]] for r in settlements],
                 [22, 16, 22, 22, 16, 16, 16, 38])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    # Re-open in memory before returning so a bad workbook never reaches the user.
    load_workbook(output, read_only=True, data_only=False).close()
    output.seek(0)
    return output
