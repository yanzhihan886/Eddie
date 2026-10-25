from __future__ import annotations

import hashlib
import io
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook

from profit_app.config import REPORT_TYPES
from profit_app.database import connect
from profit_app.importer import parse_upload, suggest_mapping
from tests.conftest import DEMO, DEMO_ORDER, TYPE_BY_FILE
from app import create_app


def _upload_and_confirm(client, store_id: str, path: Path, report_type: str):
    response = client.post("/imports/preview", data={
        "store_id": store_id,
        "report_type": report_type,
        "file": (io.BytesIO(path.read_bytes()), path.name),
    }, content_type="multipart/form-data", follow_redirects=True)
    assert response.status_code == 200
    import_path = response.request.path
    assert re.fullmatch(r"/imports/[0-9a-f]{32}", import_path)
    token = import_path.rsplit("/", 1)[1]
    parsed = parse_upload(path, path.name)
    mapping = suggest_mapping(report_type, parsed.headers)
    form = {f"map_{key}": value for key, value in mapping.items()}
    form["action"] = "check"
    checked = client.post(import_path, data=form)
    assert checked.status_code == 200
    assert "校验样本" in checked.get_data(as_text=True)
    form["action"] = "confirm"
    result = client.post(import_path, data=form, follow_redirects=True)
    assert result.status_code == 200
    assert "导入批次完成" in result.get_data(as_text=True)
    return result, parsed, mapping


def test_full_upload_dashboard_export_and_persistence(isolated):
    client = isolated["client"]
    response = client.post("/stores", data={"action": "create", "platform": "其他/演示", "name": "端到端演示店"})
    assert response.status_code == 302
    db = connect(isolated["database"])
    store = db.execute("SELECT * FROM stores WHERE name='端到端演示店'").fetchone()
    store_id = store["id"]
    db.close()

    for filename in DEMO_ORDER:
        _upload_and_confirm(client, store_id, DEMO / filename, TYPE_BY_FILE[filename])

    dashboard = client.get("/?store_id=" + store_id + "&start=2026-10-01&end=2026-10-31")
    html = dashboard.get_data(as_text=True)
    assert dashboard.status_code == 200
    assert "¥887.36" in html
    assert "¥128.00" in html
    assert "¥403.74" in html
    assert "SKU-C" in html

    for path in ("/stores", "/imports", "/costs", "/issues?status_filter=all", "/analytics?store_id=" + store_id + "&start=2026-10-01&end=2026-10-31"):
        assert client.get(path).status_code == 200

    exported = client.get(f"/reports.xlsx?store_id={store_id}&start=2026-10-01&end=2026-10-31")
    assert exported.status_code == 200
    assert exported.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    wb = load_workbook(io.BytesIO(exported.data), data_only=False)
    assert wb.sheetnames == ["经营汇总", "订单明细", "退款明细", "平台费用", "商品成本利润", "结算差异", "数据缺失与异常", "结算记录"]
    overview = wb["经营汇总"]
    totals = {overview.cell(r, 1).value: overview.cell(r, 2).value for r in range(8, overview.max_row + 1)}
    assert totals["商品销售额"] == 887.36
    assert totals["确认退款"] == 128.0
    assert totals["已知项目利润估算"] == 403.74
    assert wb["订单明细"].max_row == 12  # 11 October order lines + header; September order is outside the filter
    assert not any(cell.data_type == "f" for sheet in wb.worksheets for row in sheet.iter_rows() for cell in row)
    wb.close()

    # Recreate the app over the same DB to prove persistence survives a restart.
    app_again = create_app({"TESTING": True, "DATABASE": isolated["database"],
                            "STAGING_DIR": str(isolated["tmp_path"] / "restart-staging"),
                            "EXPORT_DIR": str(isolated["tmp_path"] / "restart-exports")})
    assert "端到端演示店" in app_again.test_client().get("/stores").get_data(as_text=True)


def test_public_demo_isolates_visitors_and_supports_upload_to_export(tmp_path):
    demo_root = tmp_path / "public-demo-sessions"
    app = create_app({
        "TESTING": True,
        "SECRET_KEY": "test-only-session-secret",
        "DEMO_MODE": True,
        "DEMO_DATA_DIR": str(demo_root),
        "DEMO_REPORTS_DIR": str(DEMO),
        "DEMO_SESSION_TTL_HOURS": 6,
        "DEMO_MAX_SESSIONS": 50,
        "MAX_CONTENT_LENGTH": 6 * 1024 * 1024,
    })
    alice = app.test_client()
    bob = app.test_client()

    health = alice.get("/health")
    assert health.status_code == 200
    assert health.json == {"status": "ok", "mode": "public-demo"}
    assert not list(demo_root.glob("*/seller_profit.sqlite3"))

    dashboard_url = "/?start=2026-10-01&end=2026-10-31"
    alice_dashboard = alice.get(dashboard_url)
    assert alice_dashboard.status_code == 200
    assert "公开体验环境" in alice_dashboard.get_data(as_text=True)
    assert "¥887.36" in alice_dashboard.get_data(as_text=True)
    assert "公开 Demo" in alice_dashboard.get_data(as_text=True)

    cookie = alice.get_cookie("session")
    session_data = app.session_interface.get_signing_serializer(app).loads(cookie.value)
    alice_db_path = demo_root / session_data["demo_workspace_id"] / "seller_profit.sqlite3"
    db = connect(alice_db_path)
    demo_store_id = db.execute("SELECT id FROM stores WHERE name='模拟配件店（虚构数据）'").fetchone()[0]
    db.close()

    alice.post("/stores", data={"action": "create", "platform": "其他/演示", "name": "仅 Alice 可见的测试店"})
    assert "仅 Alice 可见的测试店" in alice.get("/stores").get_data(as_text=True)
    assert "仅 Alice 可见的测试店" not in bob.get("/stores").get_data(as_text=True)
    assert "¥887.36" in bob.get(dashboard_url).get_data(as_text=True)

    added_order = tmp_path / "public_demo_order.csv"
    added_order.write_text(
        "订单号,下单时间,订单状态,商品编码,商品名称,数量,商品实付金额\n"
        "PUBLIC-DEMO-001,2026-10-20,已完成,SKU-A,演示充电线,1,25.00\n",
        encoding="utf-8-sig",
    )
    _upload_and_confirm(alice, demo_store_id, added_order, "orders")

    updated_dashboard = alice.get(dashboard_url)
    assert updated_dashboard.status_code == 200
    assert "¥912.36" in updated_dashboard.get_data(as_text=True)
    exported = alice.get("/reports.xlsx?start=2026-10-01&end=2026-10-31")
    assert exported.status_code == 200
    exported_workbook = load_workbook(io.BytesIO(exported.data), data_only=False)
    export_totals = {
        exported_workbook["经营汇总"].cell(row, 1).value: exported_workbook["经营汇总"].cell(row, 2).value
        for row in range(8, exported_workbook["经营汇总"].max_row + 1)
    }
    assert export_totals["商品销售额"] == 912.36
    assert "PUBLIC-DEMO-001" in [row[2].value for row in exported_workbook["订单明细"].iter_rows(min_row=2)]
    exported_workbook.close()
    assert "¥887.36" in bob.get(dashboard_url).get_data(as_text=True)
    assert list((demo_root / session_data["demo_workspace_id"] / "staging").iterdir()) == []


def test_public_demo_caps_active_sessions(tmp_path):
    app = create_app({
        "TESTING": True,
        "SECRET_KEY": "test-only-session-secret",
        "DEMO_MODE": True,
        "DEMO_DATA_DIR": str(tmp_path / "limited-demo"),
        "DEMO_REPORTS_DIR": str(DEMO),
        "DEMO_MAX_SESSIONS": 1,
    })
    first_visitor = app.test_client()
    second_visitor = app.test_client()
    assert first_visitor.get("/").status_code == 200
    assert second_visitor.get("/").status_code == 503


def test_upload_duplicate_file_is_refused(isolated):
    client = isolated["client"]
    client.post("/stores", data={"action": "create", "platform": "其他/演示", "name": "重复测试店"})
    db = connect(isolated["database"])
    store_id = db.execute("SELECT id FROM stores WHERE name='重复测试店'").fetchone()[0]
    db.close()
    path = DEMO / "orders.csv"
    _upload_and_confirm(client, store_id, path, "orders")
    response = client.post("/imports/preview", data={
        "store_id": store_id, "report_type": "orders", "file": (io.BytesIO(path.read_bytes()), path.name),
    }, content_type="multipart/form-data", follow_redirects=True)
    assert response.status_code == 200
    assert "相同文件" in response.get_data(as_text=True)
    token = response.request.path.rsplit("/", 1)[1]
    parsed = parse_upload(path, path.name)
    form = {f"map_{key}": value for key, value in suggest_mapping("orders", parsed.headers).items()}
    form["action"] = "confirm"
    rejected = client.post(f"/imports/{token}", data=form)
    assert rejected.status_code == 400
    assert "未重复写入" in rejected.get_data(as_text=True)


def test_xlsx_sheet_selection_and_import(isolated, tmp_path):
    client = isolated["client"]
    client.post("/stores", data={"action": "create", "platform": "拼多多", "name": "XLSX样例店"})
    db = connect(isolated["database"])
    store_id = db.execute("SELECT id FROM stores WHERE name='XLSX样例店'").fetchone()[0]
    db.close()

    path = tmp_path / "multi_sheet.xlsx"
    wb = Workbook()
    wb.active.title = "说明"
    wb.active.append(["这是示例工作簿，不含真实店铺数据"])
    ws = wb.create_sheet("订单明细")
    ws.append(["订单号", "下单时间", "订单状态", "商品编码", "商品名称", "数量", "商品实付金额"])
    ws.append(["X-1", "2026-10-20", "已完成", "SKU-X", "样例商品", 2, 39.99])
    wb.save(path)

    response = client.post("/imports/preview", data={
        "store_id": store_id, "report_type": "orders", "file": (io.BytesIO(path.read_bytes()), path.name),
    }, content_type="multipart/form-data", follow_redirects=True)
    assert response.status_code == 200
    token = response.request.path.rsplit("/", 1)[1]
    assert "工作表" in response.get_data(as_text=True)
    selected = client.post(f"/imports/{token}", data={"action": "sheet", "sheet_name": "订单明细"})
    assert selected.status_code == 200
    assert "订单号" in selected.get_data(as_text=True)
    mapping = suggest_mapping("orders", ["订单号", "下单时间", "订单状态", "商品编码", "商品名称", "数量", "商品实付金额"])
    form = {f"map_{key}": value for key, value in mapping.items()}
    form["action"] = "check"
    assert client.post(f"/imports/{token}", data=form).status_code == 200
    form["action"] = "confirm"
    result = client.post(f"/imports/{token}", data=form, follow_redirects=True)
    assert result.status_code == 200
    db = connect(isolated["database"])
    assert db.execute("SELECT item_amount_fen FROM orders WHERE platform_order_id='X-1'").fetchone()[0] == 3_999
    db.close()
